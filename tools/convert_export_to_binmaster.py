"""
Converts a raw PHYSICAL_INVENTORY_EXPORT_*.xlsx (from the WMS) into the
BinMaster shape this app expects, as a CSV you can paste into the
"Master Ref" Google Sheet (replacing everything below the header row).

Usage:
    pip install openpyxl --break-system-packages
    python3 convert_export_to_binmaster.py PHYSICAL_INVENTORY_EXPORT_20260909190532.xlsx BinMaster.csv

What it does (see spec §2.1):
  1. Drops rows where Location Type = CONTAINER (these are picking totes,
     Container Code is always "PICK_CONTAINER" — not a real bin).
  2. Groups the remaining rows by (Container Code, SKU Code, ERP Batch
     Number, Basic UOM).
  3. Sums Total quantity within each group (this collapses the
     Inventory Status split — UR/Block/QI — into one number, per your
     decision not to track status separately).
  4. Carries through SKU Name and the earliest Expiration Date in the
     group as read-only display context.
"""
import csv
import sys
from collections import defaultdict

import openpyxl


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    in_path, out_path = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(in_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    # Expected column order from the WMS export:
    # Owner, Owner Code, SKU Code, SKU Name, Inventory Status, ERP Batch Number,
    # Total quantity, Occupied Qty., Available Qty., Basic UOM, Warehouse area code,
    # Location Type, Container Code, Bin-location Code, Manufacturing Date, Expiration Date
    IDX_SKU, IDX_SKU_NAME, IDX_BATCH = 2, 3, 5
    IDX_QTY, IDX_UOM = 6, 9
    IDX_WH_AREA, IDX_LOC_TYPE, IDX_CONTAINER = 10, 11, 12
    IDX_EXP_DATE = 15

    groups = defaultdict(lambda: {"qty": 0.0, "sku_name": "", "wh_area": "", "exp_dates": []})

    skipped_container_rows = 0
    for r in rows:
        if r[IDX_LOC_TYPE] == "CONTAINER":
            skipped_container_rows += 1
            continue
        key = (r[IDX_CONTAINER], r[IDX_SKU], r[IDX_BATCH], r[IDX_UOM])
        g = groups[key]
        g["qty"] += r[IDX_QTY] or 0
        g["sku_name"] = r[IDX_SKU_NAME] or g["sku_name"]
        g["wh_area"] = r[IDX_WH_AREA] or g["wh_area"]
        if r[IDX_EXP_DATE]:
            g["exp_dates"].append(r[IDX_EXP_DATE])

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Bin", "Mat", "MatName", "Batch", "Qty", "UOM", "ExpirationDate", "WarehouseArea"])
        for (bin_code, sku, batch, uom), g in sorted(groups.items()):
            exp_date = min(g["exp_dates"]).strftime("%Y-%m-%d") if g["exp_dates"] else ""
            writer.writerow([bin_code, sku, g["sku_name"], batch, g["qty"], uom, exp_date, g["wh_area"]])

    print(f"Wrote {len(groups)} BinMaster rows to {out_path}")
    print(f"Skipped {skipped_container_rows} CONTAINER (PICK_CONTAINER) rows")


if __name__ == "__main__":
    main()
